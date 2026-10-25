import sys
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
from budget_draft import load_client, plan, write_workbook  # noqa: E402
from funnel import add_mom, scale_to_month, summarise  # noqa: E402

OCT_MTD = [date(2026, 10, d) for d in range(1, 27)]  # cut-off on the 26th


def test_flat_daily_value_scales_to_full_month():
    projected, factor = scale_to_month({d: Decimal("100") for d in OCT_MTD}, "2026-10")
    assert projected == Decimal("3100")
    assert factor == Decimal("3100") / Decimal("2600")


def test_weekend_heavy_days_are_weighted_by_weekday_count():
    daily = {d: Decimal("200") if d.weekday() >= 5 else Decimal("100") for d in OCT_MTD}
    projected, _ = scale_to_month(daily, "2026-10")
    # October 2026 has 9 weekend days and 22 weekdays.
    assert projected == Decimal("200") * 9 + Decimal("100") * 22


def _daily(outlet_id, platform="swiggy", orders="40", days=OCT_MTD):
    return [{
        "client": "tbs", "platform": platform, "outlet_id": outlet_id,
        "outlet_name": f"Outlet {outlet_id}", "date": d,
        "impressions": Decimal("1000"), "menu_visits": Decimal("200"), "cart": Decimal("60"),
        "orders": Decimal(orders), "gmv": Decimal("12000"), "cpc_spend": Decimal("300"),
        "cba_spend": Decimal("100"), "cpc_attributed_gmv": Decimal("900"),
        "cba_attributed_gmv": Decimal("400"), "cpc_menu_visits": Decimal("80"),
    } for d in days]


def test_summarise_ratios_and_proxy():
    [rec] = summarise(_daily("1379344"), "2026-10")
    assert rec["mtd_cutoff"] == "2026-10-26"
    assert rec["menu_conv"] == Decimal("0.2")
    assert rec["cart_conv"] == Decimal("0.3")
    assert rec["order_conv"] == Decimal("40") / Decimal("60")
    assert rec["cpc_roas"] == Decimal("3")
    assert rec["cba_roas"] == Decimal("4")
    assert rec["roas"] == Decimal("1300") / Decimal("400")
    assert rec["ad_pct_gmv"] == Decimal("400") / Decimal("12000")
    assert rec["orders_per_day_recent"] == Decimal("40")
    assert rec["cpc_clicks_proxy"] == Decimal("80") * 31


def test_mom_compares_with_previous_month():
    [rec] = summarise(_daily("1379344"), "2026-10")
    prev = {"platform": "swiggy", "outlet_id": "1379344", "orders": str(rec["orders"] / 2),
            "gmv": "", "ad_spend": "", "menu_conv": "", "cart_conv": "", "order_conv": "",
            "roas": ""}
    [rec] = add_mom([rec], [prev])
    assert rec["orders_mom"] == Decimal("1")
    assert rec["gmv_mom"] is None


@pytest.fixture
def tbs():
    return load_client("tbs")


def _funnel(outlet_id, platform, orders_per_day="40", cpc_roas="3", cba_roas="", gmv="300000"):
    return {"platform": platform, "outlet_id": outlet_id, "outlet_name": "", "gmv": gmv,
            "orders_per_day_recent": orders_per_day, "cpc_roas": cpc_roas, "cba_roas": cba_roas}


LAST_BUDGET = [
    {"platform": "swiggy", "outlet_id": "1379344", "product": "cpc", "amount": "6000"},
    {"platform": "swiggy", "outlet_id": "1379344", "product": "cba", "amount": "2000"},
    {"platform": "zomato", "outlet_id": "22734640", "product": "cpc", "amount": "2000"},
    {"platform": "zomato", "outlet_id": "22734640", "product": "cba", "amount": "500"},
    {"platform": "zomato", "outlet_id": "19356224", "product": "cpc", "amount": "1500"},
]


def test_plan_applies_rules(tbs):
    config, rules = tbs
    funnel = [
        _funnel("1379344", "swiggy", cpc_roas="3", cba_roas="5"),
        _funnel("22734640", "zomato", cpc_roas="2"),
        _funnel("22566498", "zomato", orders_per_day="35"),   # traction, no budget last month
        _funnel("22645595", "zomato", orders_per_day="10"),   # no traction
        _funnel("19356224", "zomato", orders_per_day="80"),   # excluded store
    ]
    carried, new, dropped, gmv = plan(config, rules, LAST_BUDGET, funnel)

    assert {(d["outlet_id"], d["product"]) for d in dropped} == {("22734640", "cba"), ("19356224", "cpc")}
    assert sum(r["carried_ratio"] for r in carried) == 1
    assert [r["carried_ratio"] for r in carried] == [Decimal("0.6"), Decimal("0.2"), Decimal("0.2")]
    by_key = {(r["outlet_id"], r["product"]): r for r in carried}
    assert by_key[("1379344", "cpc")]["outlet"] == "Potheri"
    assert "CPC kept" in by_key[("1379344", "cpc")]["reason"]
    assert "below break-even" in by_key[("22734640", "cpc")]["reason"]
    assert [(r["outlet_id"], r["product"]) for r in new] == [("22566498", "cpc")]
    assert gmv == Decimal("1200000")  # excluded store's GMV left out


def test_workbook_total_is_input_and_amounts_are_formulas(tbs, tmp_path):
    config, rules = tbs
    funnel = [_funnel("1379344", "swiggy"), _funnel("22566498", "zomato", orders_per_day="35")]
    carried, new, dropped, gmv = plan(config, rules, LAST_BUDGET[:2], funnel)
    path = write_workbook("tbs", "2026-11", rules, carried, new, dropped, gmv, LAST_BUDGET[:2],
                          tmp_path / "budget_draft.xlsx")
    wb = load_workbook(path)
    ws = wb["budget_draft"]
    assert ws["A2"].value == "Grand total (enter)" and ws["B2"].value is None
    assert ws["B3"].value == "=SUM(G11:G11)"
    assert ws["B4"].value == "=B2-B3"
    assert [ws.cell(row=r, column=7).value for r in (9, 10)] == ["=$B$4*F9", "=$B$4*F10"]
    assert ws.cell(row=11, column=7).value is None  # new outlet: Vicky sets the amount
    assert ws["A12"].value == "Total"
    assert wb.sheetnames == ["budget_draft", "carry_forward", "dropped"]
    assert wb["carry_forward"].max_row == 3
