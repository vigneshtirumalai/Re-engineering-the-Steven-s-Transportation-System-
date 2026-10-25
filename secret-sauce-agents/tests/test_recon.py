import sys
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import paths  # noqa: E402
from build_recon_workbook import build  # noqa: E402
from recon_schema import (  # noqa: E402
    payout_breaks, read_tidy, tie_out_breaks, to_decimal, write_tidy,
)


def _row(outlet_id, field, amount, platform="zomato", raw_label=""):
    return {
        "client": "tempt", "platform": platform, "outlet_id": outlet_id,
        "outlet_name": f"Outlet {outlet_id}", "period_start": "2026-10-12",
        "period_end": "2026-10-18", "payout_date": "2026-10-21", "field": field,
        "amount": Decimal(amount), "raw_label": raw_label or field,
        "source_file": "synthetic.xlsx", "source_row": "1",
    }


def _outlet(outlet_id, net="747.50", payout=None):
    """Synthetic outlet: 1000 sales + 50 GST collected, 302.50 of deductions."""
    rows = [
        _row(outlet_id, "gross_taxable_sales", "1000.00"),
        _row(outlet_id, "gst_9_5_collected", "50.00"),
        _row(outlet_id, "commission", "180.00"),
        _row(outlet_id, "gst_on_commission", "32.40"),
        _row(outlet_id, "ads_discount", "20.00"),
        _row(outlet_id, "tds_194o", "10.00"),
        _row(outlet_id, "tcs", "10.10"),
        _row(outlet_id, "gst_9_5_remitted", "50.00"),
        _row(outlet_id, "net_settlement", net),
    ]
    if payout is not None:
        rows.append(_row(outlet_id, "payout_amount", payout))
    return rows


def test_to_decimal_handles_platform_formats():
    assert to_decimal("1,234.50") == Decimal("1234.50")
    assert to_decimal("(12.00)") == Decimal("-12.00")
    assert to_decimal("₹ 5") == Decimal("5")
    assert to_decimal("") == Decimal("0")
    with pytest.raises(ValueError):
        to_decimal("abc")


def test_tie_out_passes_within_one_rupee():
    assert tie_out_breaks(_outlet("21306274", net="747.90")) == {}


def test_tie_out_flags_break_over_one_rupee():
    breaks = tie_out_breaks(_outlet("21306274", net="760.00"))
    assert breaks == {("zomato", "21306274"): Decimal("12.50")}


def test_payout_bridge_with_adjustment():
    rows = _outlet("21306274", payout="737.50") + _outlet("21776569", payout="747.50")
    rows.append(_row("21306274", "prior_period_adjustments", "-10.00"))
    assert payout_breaks(rows) == {}


def test_payout_bridge_flags_unexplained_difference():
    rows = _outlet("21306274", payout="700.00")
    assert payout_breaks(rows) == {"zomato": Decimal("-47.50")}


def test_unmapped_lines_never_change_totals_and_reach_workbook(tmp_path):
    rows = _outlet("21306274", payout="747.50")
    rows.append(_row("21306274", "UNMAPPED", "99.00", raw_label="Mystery fee"))
    assert tie_out_breaks(rows) == {}
    wb = load_workbook(build(rows, tmp_path / "recon.xlsx"))
    unmapped = [r for r in wb["UNMAPPED"].iter_rows(min_row=2, values_only=True)]
    assert unmapped == [("zomato", "21306274", "Mystery fee", 99, "synthetic.xlsx", "1")]


def test_workbook_grid_totals_and_keep_rates(tmp_path):
    rows = _outlet("21306274", payout="747.50") + _outlet("21776569", payout="747.50")
    wb = load_workbook(build(rows, tmp_path / "recon.xlsx"))
    assert wb.sheetnames == ["Zomato", "UNMAPPED", "Source"]
    ws = wb["Zomato"]
    header = [c.value for c in ws[2]]
    total = {h: c.value for h, c in zip(header, ws[ws.max_row])}
    assert total["Outlet"] == "Total"
    assert total["Total Gross"] == pytest.approx(2100)
    assert total["Total Deductions"] == pytest.approx(605)
    assert total["Net Settlement"] == pytest.approx(1495)
    assert total["Cash Keep Rate"] == pytest.approx(0.7475)
    assert total["Economic Keep Rate"] == pytest.approx(0.76760)
    assert list(wb["Source"].values) == [("source_file", "rows"), ("synthetic.xlsx", 20)]


def test_tidy_csv_round_trip_keeps_exact_amounts(tmp_path):
    rows = _outlet("21306274")
    path = tmp_path / "tidy.csv"
    write_tidy(rows, path)
    assert [r["amount"] for r in read_tidy(path)] == [r["amount"] for r in rows]


GOLDEN = sorted(paths.FIXTURES_DIR.glob("tempt_golden_*.xlsx"))


@pytest.mark.skipif(not GOLDEN, reason="awaiting golden workbooks in fixtures/ (see TODO.md)")
@pytest.mark.parametrize("golden", GOLDEN, ids=lambda p: p.stem)
def test_matches_golden_within_one_rupee(golden):
    pytest.skip("Zomato parser not implemented yet; written against the first real annexure")


@pytest.mark.skip(reason="awaiting Swiggy settlement files")
def test_swiggy_parser():
    pass
