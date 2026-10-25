"""Tidy settlement CSV -> outputs/<client>/<period_start>/recon.xlsx.

Usage: python scripts/build_recon_workbook.py <tidy.csv> <recon.xlsx>
"""
import sys
from collections import Counter
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent))
from recon_schema import (  # noqa: E402
    BRIDGE, CREDIT, DEDUCTIONS, REVENUE, UNMAPPED, outlet_totals, read_tidy,
)

GRID = [
    ("Revenue", list(REVENUE) + ["total_gross"]),
    ("Deductions", list(DEDUCTIONS) + ["total_deductions"]),
    ("Credit", list(CREDIT)),
    ("Bridge", list(BRIDGE)),
    ("Keep rate", ["cash_keep_rate", "economic_keep_rate"]),
]
LABELS = {
    **REVENUE, **DEDUCTIONS, **CREDIT, **BRIDGE,
    "total_gross": "Total Gross",
    "total_deductions": "Total Deductions",
    "cash_keep_rate": "Cash Keep Rate",
    "economic_keep_rate": "Economic Keep Rate",
}
BOLD = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="F0EFEC")
MONEY = "#,##0.00"
PCT = "0.0%"


def _platform_sheet(wb, platform, totals):
    ws = wb.create_sheet(platform.capitalize())
    columns = [f for _, fields in GRID for f in fields]
    group_row = ["", ""]
    for group, fields in GRID:
        group_row += [group] + [""] * (len(fields) - 1)
    ws.append(group_row)
    ws.append(["Outlet ID", "Outlet"] + [LABELS[f] for f in columns])
    for row in ws.iter_rows(min_row=1, max_row=2):
        for cell in row:
            cell.font = BOLD
            cell.fill = HEADER_FILL

    keys = sorted(k for k in totals if k[0] == platform)
    for key in keys:
        t = totals[key]
        ws.append([key[1], t["outlet_name"]] + [t[f] for f in columns])

    total_row = ["", "Total"]
    for f in columns:
        if f.endswith("keep_rate"):
            total_row.append(None)
        else:
            total_row.append(sum((totals[k][f] for k in keys), Decimal("0")))
    gts = total_row[2 + columns.index("gross_taxable_sales")]
    if gts:
        net = total_row[2 + columns.index("net_settlement")]
        credits = total_row[2 + columns.index("tds_194o")] + total_row[2 + columns.index("tcs")]
        total_row[2 + columns.index("cash_keep_rate")] = net / gts
        total_row[2 + columns.index("economic_keep_rate")] = (net + credits) / gts
    ws.append(total_row)
    for cell in ws[ws.max_row]:
        cell.font = BOLD

    for row in ws.iter_rows(min_row=3, min_col=3):
        for cell, field in zip(row, columns):
            cell.number_format = PCT if field.endswith("keep_rate") else MONEY
    ws.freeze_panes = "C3"
    ws.column_dimensions["B"].width = 22
    return ws


def build(rows, out_path):
    wb = Workbook()
    wb.remove(wb.active)
    totals = outlet_totals(rows)
    for platform in sorted({p for p, _ in totals}):
        _platform_sheet(wb, platform, totals)

    unmapped = wb.create_sheet("UNMAPPED")
    unmapped.append(["platform", "outlet_id", "raw_label", "amount", "source_file", "source_row"])
    for row in rows:
        if row["field"] == UNMAPPED:
            unmapped.append([row["platform"], row["outlet_id"], row["raw_label"],
                             row["amount"], row["source_file"], row["source_row"]])

    source = wb.create_sheet("Source")
    source.append(["source_file", "rows"])
    for name, count in sorted(Counter(r["source_file"] for r in rows).items()):
        source.append([name, count])

    for ws in (unmapped, source):
        for cell in ws[1]:
            cell.font = BOLD
            cell.fill = HEADER_FILL

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    print(build(read_tidy(sys.argv[1]), sys.argv[2]))
