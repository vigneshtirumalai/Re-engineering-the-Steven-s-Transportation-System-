"""One local page for all four agents: outputs/dashboard/index.html.

Usage: python scripts/build_dashboard.py [--today YYYY-MM-DD]

Reads only what the workflows already wrote (outputs/, logs/, TODO.md) plus today's
morning plan. Computes no business numbers itself: keep rates come from recon.xlsx, counts
from the logs. The page is self-contained and stays on this machine (outputs/ is gitignored).
"""
import argparse
import csv
import html
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import morning_plan  # noqa: E402
import paths  # noqa: E402

RECON_PERIOD = re.compile(r"^\d{4}-\d{2}-\d{2}$")
BUDGET_PERIOD = re.compile(r"^\d{4}-\d{2}$")
# Trust bars from the build plan: max edits on numbers, max edits on text, max minutes.
TRUST_BARS = {"recon": (0, 1, 15), "digest": (0, 1, 10), "adbudget": (1, 2, 60)}
GRADUATE_WEEKS = 4
MIN_SCORED_FOR_BAR = 8


# ---------- collection (pure reads) ----------

def _read_csv(path):
    if not path.exists():
        return []
    with open(path, newline="") as f:
        return list(csv.DictReader(f))


def _text(path):
    return path.read_text() if path.exists() else None


def _clients(outputs_dir):
    return sorted(p.name for p in outputs_dir.iterdir()
                  if p.is_dir() and p.name not in ("brief", "dashboard")) if outputs_dir.exists() else []


def latest_brief(outputs_dir):
    briefs = sorted((outputs_dir / "brief").glob("*.md"))
    if not briefs:
        return None
    text = briefs[-1].read_text()
    return {"date": briefs[-1].stem, "path": briefs[-1], "text": text,
            "failed": "FAILED" in text.splitlines()[0] if text else True}


def keep_rates(workbook_path):
    """Per-outlet rows from recon.xlsx platform sheets, as the workbook states them."""
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    rows = []
    for ws in wb.worksheets:
        if ws.title in ("UNMAPPED", "Source"):
            continue
        values = list(ws.iter_rows(values_only=True))
        if len(values) < 3:
            continue
        header = list(values[1])
        col = {name: header.index(name) for name in
               ("Outlet ID", "Outlet", "Net Settlement", "Cash Keep Rate", "Economic Keep Rate")
               if name in header}
        for v in values[2:]:
            rows.append({"platform": ws.title, **{k: v[i] for k, i in col.items()}})
    unmapped = wb["UNMAPPED"].max_row - 1 if "UNMAPPED" in wb.sheetnames else 0
    wb.close()
    return rows, unmapped


def recon_status(outputs_dir):
    out = []
    for client in _clients(outputs_dir):
        periods = sorted(p for p in (outputs_dir / client).iterdir()
                         if p.is_dir() and RECON_PERIOD.match(p.name))
        if not periods:
            continue
        p = periods[-1]
        checker = _text(p / "checker.md")
        verdict = checker.strip().splitlines()[0].strip() if checker and checker.strip() else None
        rec = {"client": client, "period": p.name, "dir": p, "verdict": verdict,
               "checker": checker, "summary": _text(p / "summary.md"),
               "missing": _text(p / "MISSING.md"), "workbook": None, "rates": [], "unmapped": 0}
        if (p / "recon.xlsx").exists():
            rec["workbook"] = p / "recon.xlsx"
            rec["rates"], rec["unmapped"] = keep_rates(p / "recon.xlsx")
        out.append(rec)
    return out


def digest_status(outputs_dir):
    out = []
    for client in _clients(outputs_dir):
        drafts = sorted((outputs_dir / client / "digests").glob("*_whatsapp.txt"))
        if not drafts:
            continue
        d = drafts[-1]
        day = d.name.removesuffix("_whatsapp.txt")
        verify = _text(d.with_name(f"{day}_verify.txt"))
        out.append({"client": client, "date": day, "path": d, "draft": d.read_text(),
                    "verify": verify.strip() if verify and verify.strip() else None})
    return out


def budget_status(outputs_dir):
    out = []
    for client in _clients(outputs_dir):
        months = sorted(p for p in (outputs_dir / client).iterdir()
                        if p.is_dir() and BUDGET_PERIOD.match(p.name))
        if not months:
            continue
        m = months[-1]
        out.append({"client": client, "month": m.name, "dir": m,
                    "draft": m / "budget_draft.xlsx" if (m / "budget_draft.xlsx").exists() else None,
                    "memo": _text(m / "memo.md"), "missing": _text(m / "MISSING.md")})
    return out


def prediction_stats(logs_dir):
    preds = _read_csv(logs_dir / "predictions.csv")
    counts = Counter(p.get("scored") or "open" for p in preds)
    decided = counts["hit"] + counts["miss"]
    scored = decided + counts["unclear"]
    return {"total": len(preds), "open": counts["open"], "hit": counts["hit"],
            "miss": counts["miss"], "unclear": counts["unclear"], "scored": scored,
            "hit_rate": counts["hit"] / decided if decided and scored >= MIN_SCORED_FOR_BAR else None}


def _row_passes(row, bars):
    max_num, max_text, max_min = bars
    try:
        return (int(row["edits_numbers"]) <= max_num and int(row["edits_text"]) <= max_text
                and int(row["minutes_spent"]) <= max_min)
    except (KeyError, ValueError):
        return False


def trust_streaks(logs_dir):
    """Consecutive passing ISO weeks per workflow, counted back from the latest logged week."""
    weeks = defaultdict(lambda: defaultdict(list))
    for row in _read_csv(logs_dir / "review_log.csv"):
        if row.get("workflow") in TRUST_BARS:
            y, w, _ = date.fromisoformat(row["date"]).isocalendar()
            weeks[row["workflow"]][(y, w)].append(row)
    out = {}
    for wf, bars in TRUST_BARS.items():
        streak, prev = 0, None
        for key in sorted(weeks[wf], reverse=True):
            contiguous = prev is None or date.fromisocalendar(*key, 1).toordinal() == \
                date.fromisocalendar(*prev, 1).toordinal() - 7
            if not contiguous or not all(_row_passes(r, bars) for r in weeks[wf][key]):
                break
            streak, prev = streak + 1, key
        out[wf] = {"streak": streak, "reviews": sum(len(v) for v in weeks[wf].values())}
    return out


def open_todos(todo_path):
    text = _text(todo_path) or ""
    section, items = "", []
    for line in text.splitlines():
        if line.startswith("## "):
            section = line[3:].strip()
        elif line.startswith("- [ ] "):
            items.append((section, line[6:].strip()))
    return items


def all_missing(outputs_dir):
    found = []
    for p in sorted(outputs_dir.glob("*/*/MISSING.md")):
        found.append({"where": f"{p.parent.parent.name}/{p.parent.name}", "path": p, "text": p.read_text()})
    return found


def _ready(script):
    return "NotImplementedError" not in (paths.REPO_ROOT / "scripts" / script).read_text()


def agent_cards(plan, brief, recon, digests, budgets, agents_dir):
    def built(name):
        return (agents_dir / f"{name}.md").exists()

    cards = []
    # Payout recon
    if not built("payout-recon"):
        st, note = "idle", "Not built"
    elif not _ready("parse_zomato_settlement.py"):
        st, note = "blocked", "Zomato parser waits for settlement files"
    elif any(r["verdict"] == "FAIL" for r in recon):
        st, note = "failed", "Checker FAIL on latest week"
    elif any(r["missing"] for r in recon):
        st, note = "attention", "Inputs missing"
    else:
        st, note = ("ok", "Checker PASS") if recon else ("idle", "No weeks run yet")
    cards.append({"name": "Payout recon", "command": "/recon <client> <period_start>",
                  "status": st, "note": note,
                  "last": max((r["period"] for r in recon), default=None)})
    # Client digest
    if not built("client-digest"):
        st, note = "idle", "Not built (Prompt 2)"
    elif any(d["verify"] for d in digests):
        st, note = "attention", "Drafts have numbers to verify"
    else:
        st, note = ("ok", "Drafts ready") if digests else ("idle", "No drafts yet")
    cards.append({"name": "Client digest", "command": "/digest all", "status": st, "note": note,
                  "last": max((d["date"] for d in digests), default=None)})
    # Ad budget
    step = next((s for s in plan["steps"] if s["step"] == "adbudget"), None)
    if not built("ad-budget"):
        st, note = "idle", "Not built"
    elif not _ready("parse_ads_exports.py"):
        st, note = "blocked", "Export parser waits for TBS exports"
    elif step and step.get("missing"):
        st, note = "attention", "Exports missing"
    else:
        st, note = ("ok", "Draft ready") if any(b["draft"] for b in budgets) else ("idle", step["reason"] if step else "Not due")
    cards.append({"name": "Ad budget", "command": "/adbudget <client> <YYYY-MM>", "status": st,
                  "note": note, "last": max((b["month"] for b in budgets), default=None)})
    # Morning orchestrator
    if brief is None:
        st, note = "idle", "No run yet"
    elif brief["failed"]:
        st, note = "failed", "Last run FAILED"
    else:
        st, note = "ok", "Last run completed"
    cards.append({"name": "Morning run", "command": "bash scripts/run_morning.sh", "status": st,
                  "note": note, "last": brief["date"] if brief else None})
    return cards


def collect(today, outputs_dir=None, logs_dir=None, inbox_dir=None, todo_path=None, agents_dir=None):
    outputs_dir = outputs_dir or paths.OUTPUTS_DIR
    logs_dir = logs_dir or paths.LOGS_DIR
    agents_dir = agents_dir or morning_plan.AGENTS_DIR
    plan = morning_plan.plan(today, inbox_dir or paths.INBOX_DIR, outputs_dir,
                             logs_dir / "last_run.json")
    brief = latest_brief(outputs_dir)
    recon, digests, budgets = recon_status(outputs_dir), digest_status(outputs_dir), budget_status(outputs_dir)
    return {
        "today": today, "plan": plan, "brief": brief, "recon": recon, "digests": digests,
        "budgets": budgets, "predictions": prediction_stats(logs_dir),
        "trust": trust_streaks(logs_dir), "missing": all_missing(outputs_dir),
        "todos": open_todos(todo_path or paths.REPO_ROOT / "TODO.md"),
        "agents": agent_cards(plan, brief, recon, digests, budgets, agents_dir),
    }


# ---------- rendering ----------

STATUS = {  # icon + label always accompany the color
    "ok": ("✓", "OK"), "attention": ("!", "Needs you"), "blocked": ("■", "Blocked"),
    "failed": ("✕", "Failed"), "idle": ("–", "Idle"),
}
e = html.escape


def chip(status, label=None):
    icon, text = STATUS[status]
    return f'<span class="chip {status}"><span aria-hidden="true">{icon}</span>{e(label or text)}</span>'


def link(path, label=None):
    return f'<a href="{e(Path(path).resolve().as_uri())}">{e(label or Path(path).name)}</a>'


def copy_btn(text, label="Copy"):
    return f'<button class="copy" data-copy="{e(text)}">{e(label)}</button>'


def md(text):
    """Small markdown subset: headings, bullets, numbered items, code fences, bold, code."""
    out, in_list, in_code = [], None, False
    for raw in (text or "").splitlines():
        if raw.startswith("```"):
            if in_list:
                out.append(f"</{in_list}>")
                in_list = None
            out.append("</pre>" if in_code else "<pre>")
            in_code = not in_code
            continue
        if in_code:
            out.append(e(raw))
            continue
        line = e(raw)
        line = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", line)
        line = re.sub(r"`(.+?)`", r"<code>\1</code>", line)
        kind = "ul" if re.match(r"^\s*- ", raw) else "ol" if re.match(r"^\s*\d+\. ", raw) else None
        if in_list and kind != in_list:
            out.append(f"</{in_list}>")
            in_list = None
        if kind:
            if not in_list:
                out.append(f"<{kind}>")
                in_list = kind
            out.append("<li>" + re.sub(r"^\s*(- |\d+\. )", "", line) + "</li>")
        elif raw.startswith("#"):
            level = min(len(raw) - len(raw.lstrip("#")) + 2, 5)
            out.append(f"<h{level}>{line.lstrip('#').strip()}</h{level}>")
        elif line.strip():
            out.append(f"<p>{line}</p>")
    if in_list:
        out.append(f"</{in_list}>")
    if in_code:
        out.append("</pre>")
    return "\n".join(out)


def decide_today(brief_text):
    """Bullets under the brief's 'Decide today' heading."""
    items, inside = [], False
    for line in brief_text.splitlines():
        if line.lstrip("#").strip().lower().startswith("decide today") and line.startswith("#"):
            inside = True
        elif line.startswith("#"):
            inside = False
        elif inside and line.strip().startswith("- "):
            items.append(line.strip()[2:])
    return items


def _pct(v):
    return f"{v:.1%}" if isinstance(v, (int, float)) else "–"


def _money(v):
    return f"₹{v:,.0f}" if isinstance(v, (int, float)) else "–"


def render(d):
    today = d["today"]
    parts = []

    # Decisions first: the brief's "Decide today" section
    decide = decide_today(d["brief"]["text"]) if d["brief"] else []
    parts.append('<section><h2>Decide today</h2><div class="decide">'
                 + ("<ul>" + "".join(f"<li>{md(x)[3:-4]}</li>" for x in decide) + "</ul>" if decide
                    else '<p class="empty">Nothing waiting on you in the latest brief.</p>')
                 + "</div></section>")

    # Agent status row
    cards = "".join(
        f'<article class="card"><header><h3>{e(a["name"])}</h3>{chip(a["status"])}</header>'
        f'<p class="note">{e(a["note"])}</p>'
        f'<p class="meta">Last output: {e(a["last"] or "none")}</p>'
        f'<p class="cmd"><code>{e(a["command"])}</code>{copy_btn(a["command"])}</p></article>'
        for a in d["agents"])
    parts.append(f'<section><h2>Agents</h2><div class="grid cards">{cards}</div></section>')

    # Today's plan + decisions from the brief
    rows = "".join(
        f'<tr><td>{e(s["step"])}</td><td>{e(s.get("client", ""))} {e(s.get("platform", s.get("budget_month", "")))}</td>'
        f'<td>{chip("ok", "Runs") if s["run"] else chip("idle", "Skipped")}</td><td>{e(s["reason"])}</td></tr>'
        for s in d["plan"]["steps"])
    brief = d["brief"]
    brief_html = (f'<p class="meta">{link(brief["path"], "Brief " + brief["date"])}'
                  f'{" " + chip("failed", "Run failed") if brief["failed"] else ""}</p>'
                  f'<div class="md">{md(brief["text"])}</div>') if brief else '<p class="empty">No brief yet. The first 06:00 run writes one.</p>'
    parts.append(
        f'<section class="grid two"><div><h2>Today\'s plan</h2><table><thead><tr><th>Step</th><th>For</th>'
        f'<th>Status</th><th>Why</th></tr></thead><tbody>{rows}</tbody></table></div>'
        f'<div><h2>Latest brief</h2>{brief_html}</div></section>')

    # WhatsApp drafts
    if d["digests"]:
        drafts = "".join(
            f'<article class="card"><header><h3>{e(x["client"].upper())} · {e(x["date"])}</h3>'
            f'{chip("attention", "Verify first") if x["verify"] else chip("ok", "Ready")}</header>'
            f'<pre class="draft">{e(x["draft"])}</pre>'
            + (f'<p class="warn"><strong>Verify before sending:</strong> {e(x["verify"])}</p>' if x["verify"] else "")
            + f'<p>{copy_btn(x["draft"], "Copy draft")} {link(x["path"])}</p></article>'
            for x in d["digests"])
    else:
        drafts = '<p class="empty">No drafts yet. The client digest agent (Prompt 2) is not built.</p>'
    parts.append(f'<section><h2>WhatsApp drafts</h2><div class="grid cards">{drafts}</div></section>')

    # Recon
    if d["recon"]:
        blocks = []
        for r in d["recon"]:
            verdict = chip("ok", "Checker PASS") if r["verdict"] == "PASS" else \
                chip("failed", "Checker FAIL") if r["verdict"] == "FAIL" else chip("idle", "Not checked")
            rate_rows = "".join(
                f'<tr><td>{e(str(x.get("Outlet ID") or ""))}</td><td>{e(str(x.get("Outlet") or ""))}</td>'
                f'<td>{e(x["platform"])}</td><td class="num">{_money(x.get("Net Settlement"))}</td>'
                f'<td class="num">{_pct(x.get("Cash Keep Rate"))}</td><td class="num">{_pct(x.get("Economic Keep Rate"))}</td></tr>'
                for x in r["rates"])
            table = (f'<table><thead><tr><th>Outlet ID</th><th>Outlet</th><th>Platform</th><th class="num">Net settlement</th>'
                     f'<th class="num">Cash keep</th><th class="num">Economic keep</th></tr></thead><tbody>{rate_rows}</tbody></table>'
                     if rate_rows else "")
            review = f"python scripts/log_review.py recon {r['client']} 0 0 10 --output {r['dir']}"
            blocks.append(
                f'<article class="card wide"><header><h3>{e(r["client"].upper())} · week of {e(r["period"])}</h3>'
                f'<span class="chips">{verdict}'
                + (chip("attention", f'{r["unmapped"]} unmapped') if r["unmapped"] else "") + '</span></header>'
                + (f'<div class="md">{md(r["summary"])}</div>' if r["summary"] else "")
                + (f'<div class="md warn">{md(r["missing"])}</div>' if r["missing"] else "")
                + table
                + f'<p>{link(r["workbook"]) if r["workbook"] else ""} {link(r["dir"] / "checker.md") if r["checker"] else ""}'
                f' {copy_btn(review, "Copy review log command")}</p></article>')
        recon_html = "".join(blocks)
    else:
        recon_html = '<p class="empty">No recon weeks yet. Waiting on Tempt Zomato settlement files.</p>'
    parts.append(f'<section><h2>Payout recon</h2>{recon_html}</section>')

    # Ad budget
    if d["budgets"]:
        recon_like = []
        for b in d["budgets"]:
            recon_like.append(
                f'<article class="card wide"><header><h3>{e(b["client"].upper())} · budget for {e(b["month"])}</h3>'
                f'{chip("ok", "Draft ready") if b["draft"] else chip("attention", "No draft")}</header>'
                + (f'<div class="md">{md(b["memo"])}</div>' if b["memo"] else "")
                + (f'<div class="md warn">{md(b["missing"])}</div>' if b["missing"] else "")
                + f'<p>{link(b["draft"]) if b["draft"] else ""}</p></article>')
        budget_html = "".join(recon_like)
    else:
        budget_html = '<p class="empty">No budget drafts yet. Runs from the 27th once TBS exports are in.</p>'
    parts.append(f'<section><h2>Ad budget</h2>{budget_html}</section>')

    # Trust + predictions
    p = d["predictions"]
    tiles = "".join(
        f'<div class="tile"><span class="big">{v}</span><span class="label">{label}</span></div>'
        for v, label in [(p["open"], "Open"), (p["hit"], "Hit"), (p["miss"], "Miss"), (p["unclear"], "Unclear")])
    rate = (f'<p>Hit rate: <strong>{p["hit_rate"]:.0%}</strong> of {p["hit"] + p["miss"]} decided</p>'
            if p["hit_rate"] is not None else
            f'<p class="meta">Hit rate shown once {MIN_SCORED_FOR_BAR}+ predictions are scored ({p["scored"]} so far).</p>')
    names = {"recon": "Payout recon", "digest": "Client digest", "adbudget": "Ad budget"}
    trust_rows = "".join(
        f'<tr><td>{names[wf]}</td><td><span class="dots" aria-label="{t["streak"]} of {GRADUATE_WEEKS} weeks">'
        + "".join(f'<span class="dot{" on" if i < t["streak"] else ""}"></span>' for i in range(GRADUATE_WEEKS))
        + f'</span> {min(t["streak"], GRADUATE_WEEKS)}/{GRADUATE_WEEKS} weeks</td>'
        f'<td>{chip("ok", "Exception-only") if t["streak"] >= GRADUATE_WEEKS else chip("idle", "Full review")}</td>'
        f'<td class="num">{t["reviews"]}</td></tr>'
        for wf, t in d["trust"].items())
    parts.append(
        f'<section class="grid two"><div><h2>Trust</h2><table><thead><tr><th>Agent</th><th>Streak</th><th>Review mode</th>'
        f'<th class="num">Reviews logged</th></tr></thead><tbody>{trust_rows}</tbody></table>'
        f'<p class="meta">4 passing weeks in a row moves an agent to exception-only review; one miss resets it.</p>'
        f'<p class="cmd"><code>python scripts/log_review.py &lt;workflow&gt; &lt;client&gt; &lt;edits_numbers&gt; &lt;edits_text&gt; &lt;minutes&gt;</code></p></div>'
        f'<div><h2>Predictions</h2><div class="tiles">{tiles}</div>{rate}</div></section>')

    # Missing + to-do
    miss = "".join(f'<li><strong>{e(m["where"])}</strong>: {link(m["path"])}<div class="md">{md(m["text"])}</div></li>'
                   for m in d["missing"])
    for s in d["plan"]["steps"]:
        for item in s.get("missing", []):
            miss += f'<li><strong>{e(s["step"])} {e(s.get("client", ""))}</strong>: {e(item)}</li>'
    todo_groups = defaultdict(list)
    for section, item in d["todos"]:
        todo_groups[section].append(item)
    todos = "".join(f'<h4>{e(sec)}</h4><ul>' + "".join(f"<li>{md(i)[3:-4]}</li>" for i in items) + "</ul>"
                    for sec, items in todo_groups.items())
    parts.append(
        f'<section class="grid two"><div><h2>Missing inputs</h2>'
        f'{"<ul>" + miss + "</ul>" if miss else "<p class=empty>Nothing missing for today&#39;s steps.</p>"}</div>'
        f'<div><h2>To do</h2>{todos or "<p class=empty>All done.</p>"}</div></section>')

    generated = datetime.now(ZoneInfo("Asia/Kolkata")).strftime("%Y-%m-%d %H:%M IST")
    return TEMPLATE.replace("{{TODAY}}", e(today.isoformat())).replace("{{GENERATED}}", e(generated)) \
        .replace("{{BODY}}", "\n".join(parts))


TEMPLATE = """<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Secret Sauce Control Room</title>
<style>
:root {
  --bg: #f6f5f1; --surface: #fcfcfb; --ink: #0b0b0b; --ink-2: #52514e; --ink-3: #7a7974;
  --line: #e1e0d9; --accent: #2a78d6;
  --good: #0ca30c; --warn: #fab219; --serious: #ec835a; --critical: #d03b3b;
  --good-text: #006300; --warn-bg: #f9eeda;
}
@media (prefers-color-scheme: dark) {
  :root:not([data-theme="light"]) {
    --bg: #121211; --surface: #1a1a19; --ink: #f0efec; --ink-2: #c3c2b7; --ink-3: #8d8c85;
    --line: #2c2c2a; --accent: #3987e5; --good-text: #0ca30c; --warn-bg: #3a2f17;
  }
}
:root[data-theme="dark"] {
  --bg: #121211; --surface: #1a1a19; --ink: #f0efec; --ink-2: #c3c2b7; --ink-3: #8d8c85;
  --line: #2c2c2a; --accent: #3987e5; --good-text: #0ca30c; --warn-bg: #3a2f17;
}
* { box-sizing: border-box; }
body { margin: 0; background: var(--bg); color: var(--ink);
  font: 15px/1.5 -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif; }
header.top { display: flex; justify-content: space-between; align-items: baseline; flex-wrap: wrap;
  gap: 8px; padding: 20px 24px 8px; max-width: 1280px; margin: 0 auto; }
header.top h1 { margin: 0; font-size: 22px; }
header.top .meta { color: var(--ink-2); }
main { max-width: 1280px; margin: 0 auto; padding: 0 24px 48px; }
section { margin-top: 28px; }
h2 { font-size: 13px; text-transform: uppercase; letter-spacing: .06em; color: var(--ink-2); margin: 0 0 10px; }
h3 { font-size: 15px; margin: 0; }
h4 { margin: 12px 0 4px; font-size: 13px; color: var(--ink-2); }
.grid { display: grid; gap: 16px; }
.grid > * { min-width: 0; }
.chips { display: inline-flex; gap: 6px; flex-wrap: wrap; }
.decide { background: var(--surface); border: 1px solid var(--line); border-left: 4px solid var(--accent); border-radius: 10px; padding: 12px 16px; }
.decide ul { margin: 4px 0; padding-left: 20px; }
.cards { grid-template-columns: repeat(auto-fill, minmax(260px, 1fr)); }
.two { grid-template-columns: repeat(auto-fit, minmax(380px, 1fr)); gap: 24px; }
.card, .two > div { background: var(--surface); border: 1px solid var(--line); border-radius: 10px; padding: 14px 16px; }
.card.wide { margin-bottom: 16px; }
.card header { display: flex; justify-content: space-between; align-items: center; gap: 8px; flex-wrap: wrap; margin-bottom: 6px; }
.note { margin: 4px 0; }
.meta { color: var(--ink-3); font-size: 13px; margin: 4px 0; }
.empty { color: var(--ink-3); font-style: italic; }
.chip { display: inline-flex; align-items: center; gap: 5px; font-size: 12px; font-weight: 600;
  padding: 2px 9px; border-radius: 999px; border: 1px solid currentColor; white-space: nowrap; }
.chip.ok { color: var(--good-text); }
.chip.attention { color: var(--ink); background: var(--warn-bg); border-color: var(--warn); }
.chip.blocked { color: var(--ink); border-color: var(--serious); box-shadow: inset 3px 0 0 var(--serious); }
.chip.failed { color: #fff; background: var(--critical); border-color: var(--critical); }
.chip.idle { color: var(--ink-3); }
table { width: 100%; border-collapse: collapse; font-size: 14px; margin: 8px 0; }
th, td { text-align: left; padding: 6px 8px; border-bottom: 1px solid var(--line); vertical-align: top; }
th { color: var(--ink-2); font-weight: 600; font-size: 12px; }
.num { text-align: right; font-variant-numeric: tabular-nums; }
code { font: 12.5px ui-monospace, SFMono-Regular, Menlo, monospace; background: var(--bg); padding: 1px 5px; border-radius: 4px; }
pre { font: 12.5px/1.5 ui-monospace, SFMono-Regular, Menlo, monospace; white-space: pre-wrap; background: var(--bg);
  padding: 10px; border-radius: 6px; margin: 8px 0; overflow-wrap: anywhere; }
.cmd { display: flex; gap: 8px; align-items: center; flex-wrap: wrap; }
button.copy { font: inherit; font-size: 12px; padding: 3px 10px; border-radius: 6px; cursor: pointer;
  border: 1px solid var(--line); background: var(--surface); color: var(--ink); }
button.copy:hover { border-color: var(--accent); }
a { color: var(--accent); }
.warn { background: var(--warn-bg); border-radius: 6px; padding: 6px 10px; }
.md p { margin: 6px 0; } .md ul, .md ol { margin: 6px 0; padding-left: 20px; }
.md h3, .md h4, .md h5 { margin: 10px 0 4px; font-size: 14px; }
.tiles { display: grid; grid-template-columns: repeat(4, 1fr); gap: 10px; }
.tile { border: 1px solid var(--line); border-radius: 8px; padding: 10px; display: flex; flex-direction: column; }
.tile .big { font-size: 26px; font-weight: 650; font-variant-numeric: tabular-nums; }
.tile .label { color: var(--ink-2); font-size: 12px; }
.dots { display: inline-flex; gap: 4px; vertical-align: middle; }
.dot { width: 10px; height: 10px; border-radius: 50%; border: 1.5px solid var(--ink-3); }
.dot.on { background: var(--good); border-color: var(--good); }
.theme { font: inherit; font-size: 12px; padding: 3px 10px; border-radius: 6px; border: 1px solid var(--line);
  background: var(--surface); color: var(--ink); cursor: pointer; }
@media print {
  :root, :root[data-theme="dark"] { --bg: #fff; --surface: #fff; --ink: #0b0b0b; --ink-2: #52514e;
    --ink-3: #6b6a65; --line: #d6d5ce; --accent: #1f5fae; --good-text: #006300; --warn-bg: #f9eeda; }
  @page { size: A4; margin: 12mm; }
  body { font-size: 12px; }
  button.copy, .theme { display: none; }
  .card, .decide, .two > div, tr { break-inside: avoid; }
  .cards { grid-template-columns: repeat(2, 1fr); }
  .two { grid-template-columns: 1fr; }
  a { color: inherit; text-decoration: none; }
}
@media (max-width: 640px) {
  header.top, main { padding-left: 16px; padding-right: 16px; }
  .two { grid-template-columns: 1fr; }
  .tiles { grid-template-columns: repeat(2, 1fr); }
  table { display: block; overflow-x: auto; }
}
</style>
</head>
<body>
<header class="top">
  <h1>Secret Sauce Control Room</h1>
  <span class="meta">{{TODAY}} · refreshed {{GENERATED}} <button class="theme" id="theme">Theme</button></span>
</header>
<main>
{{BODY}}
</main>
<script>
document.addEventListener("click", async (ev) => {
  const b = ev.target.closest("button.copy");
  if (!b) return;
  try { await navigator.clipboard.writeText(b.dataset.copy); b.textContent = "Copied"; }
  catch { b.textContent = "Select and copy"; }
  setTimeout(() => (b.textContent = b.dataset.label || "Copy"), 1500);
});
document.querySelectorAll("button.copy").forEach((b) => (b.dataset.label = b.textContent));
const root = document.documentElement;
try { const t = localStorage.getItem("ss-theme"); if (t) root.dataset.theme = t; } catch {}
document.getElementById("theme").addEventListener("click", () => {
  const dark = root.dataset.theme ? root.dataset.theme === "dark"
    : matchMedia("(prefers-color-scheme: dark)").matches;
  root.dataset.theme = dark ? "light" : "dark";
  try { localStorage.setItem("ss-theme", root.dataset.theme); } catch {}
});
</script>
</body>
</html>
"""


def build(today, out_path=None, **dirs):
    out_path = Path(out_path or paths.OUTPUTS_DIR / "dashboard" / "index.html")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(render(collect(today, **dirs)))
    return out_path


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--today", type=date.fromisoformat,
                        default=datetime.now(ZoneInfo("Asia/Kolkata")).date())
    args = parser.parse_args()
    print(build(args.today))
