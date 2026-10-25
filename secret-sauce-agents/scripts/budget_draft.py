"""Apply clients/<client>/budget_rules.yaml to last month's budget and this month's funnel.

Usage: python scripts/budget_draft.py <client> <budget_month YYYY-MM> <last_budget.csv> <funnel.csv> <budget_draft.xlsx>

last_budget.csv columns: platform, outlet_id, product, amount.
funnel.csv: output of scripts/funnel.py for the month before budget_month.

The grand total is an input cell Vicky fills in; every carried amount is a formula off it.
"""
import csv
import sys
from decimal import Decimal
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent))
from paths import CLIENTS_DIR  # noqa: E402
from recon_schema import to_decimal  # noqa: E402

ID_FIELD = {"zomato": "zomato_res_id", "swiggy": "swiggy_id"}
BOLD = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="F0EFEC")
INPUT_FILL = PatternFill("solid", fgColor="F9EEDA")
RED_FILL = PatternFill("solid", fgColor="FBEBEB")
HEADERS = ["platform", "outlet_id", "outlet", "product", "last_month_amount",
           "carried_ratio", "proposed_amount", "roas_last_month", "reason"]


def load_client(client):
    config = yaml.safe_load((CLIENTS_DIR / client / "config.yaml").read_text())
    rules = yaml.safe_load((CLIENTS_DIR / client / "budget_rules.yaml").read_text())
    return config, rules


def _dec(value):
    return None if value in (None, "", "None") else to_decimal(value)


def plan(config, rules, last_budget, funnel):
    """Returns (carried rows, new-outlet rows, dropped rows, projected GMV). Pure: no I/O."""
    excluded = {(p, str(e["id"])) for p, items in (config.get("excluded_ids") or {}).items()
                for e in items}
    names = {}
    for o in config["outlets"]:
        for p, field in ID_FIELD.items():
            if o.get(field) not in (None, "TODO"):
                names[(p, str(o[field]))] = o["name"]
    perf = {(f["platform"], str(f["outlet_id"])): f for f in funnel}

    kept, dropped = [], []
    for b in last_budget:
        key = (b["platform"], str(b["outlet_id"]))
        amount = to_decimal(b["amount"])
        if key in excluded:
            dropped.append({**b, "reason": "excluded ID per config"})
        elif b["product"] not in rules["products"].get(b["platform"], []):
            dropped.append({**b, "reason": f"{b['product']} not offered on {b['platform']}"})
        else:
            kept.append({**b, "outlet_id": key[1], "amount": amount})

    pool = sum((b["amount"] for b in kept), Decimal("0"))
    break_even = 1 / Decimal(str(rules["margin"]["contribution_margin"]))
    carried = []
    for b in kept:
        key = (b["platform"], b["outlet_id"])
        f = perf.get(key, {})
        roas = _dec(f.get(f"{b['product']}_roas"))
        reasons = ["carried forward"]
        if key not in names:
            reasons.append("outlet not in config")
        if not f:
            reasons.append("no funnel data this month")
        if roas is not None and roas < break_even:
            reasons.append(f"ROAS below break-even {break_even:.1f}x")
        if b["product"] == "cpc" and rules["cpc_rule"]["never_cut_cpc_to_fund_cba"]:
            cba_roas = _dec(f.get("cba_roas"))
            if roas is not None and cba_roas is not None and cba_roas > roas:
                reasons.append("CBA ROAS higher; CPC kept (visibility rule) - trade-off to review")
        carried.append({
            "platform": b["platform"], "outlet_id": b["outlet_id"],
            "outlet": names.get(key, f.get("outlet_name", "")), "product": b["product"],
            "last_month_amount": b["amount"],
            "carried_ratio": b["amount"] / pool if pool else Decimal("0"),
            "roas_last_month": roas, "reason": "; ".join(reasons),
        })

    budgeted = {(b["platform"], b["outlet_id"]) for b in kept}
    threshold = Decimal(str(rules["new_outlets"]["min_orders_per_day"]))
    new = []
    for key, f in sorted(perf.items()):
        rate = _dec(f.get("orders_per_day_recent"))
        if key in budgeted or key in excluded or rate is None or rate < threshold:
            continue
        for product in rules["products"].get(key[0], []):
            new.append({
                "platform": key[0], "outlet_id": key[1],
                "outlet": names.get(key, f.get("outlet_name", "")), "product": product,
                "last_month_amount": None, "carried_ratio": None,
                "roas_last_month": _dec(f.get(f"{product}_roas")),
                "reason": f"new outlet with traction: {rate:.0f} orders/day over last "
                          f"{rules['new_outlets']['lookback_days']} days; set amount",
            })

    override = rules["spend_cap"].get("projected_gmv_override")
    projected_gmv = (to_decimal(override) if override is not None else
                     sum((to_decimal(f["gmv"]) for k, f in perf.items() if k not in excluded),
                         Decimal("0")))
    return carried, new, dropped, projected_gmv


def write_workbook(client, budget_month, rules, carried, new, dropped, projected_gmv,
                   last_budget, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "budget_draft"
    cap = rules["spend_cap"]["max_share_of_gmv"]
    first = 9
    last_carried = first + len(carried) - 1
    first_new = last_carried + 1
    last_new = first_new + len(new) - 1
    new_sum = f"SUM(G{first_new}:G{last_new})" if new else "0"

    ws.append([f"{client.upper()} ad budget draft for {budget_month}"])
    ws.append(["Grand total (enter)", None])
    ws.append(["New outlet lines", f"={new_sum}"])
    ws.append(["Carried pool", "=B2-B3"])
    ws.append(["Projected GMV", projected_gmv])
    ws.append(["Ad spend % of projected GMV", '=IF(OR(B5=0,B2=""),"",B2/B5)'])
    ws.append([f"Cap at {cap:.1%} of projected GMV", f"=B5*{cap}"])
    ws.append(HEADERS)
    ws["A1"].font = BOLD
    ws["B2"].fill = INPUT_FILL
    ws["B6"].number_format = "0.0%"
    ws.conditional_formatting.add("B6", CellIsRule(operator="greaterThan",
                                                   formula=[str(cap)], fill=RED_FILL))
    for cell in ws[8]:
        cell.font = BOLD
        cell.fill = HEADER_FILL

    for i, r in enumerate(carried):
        row = first + i
        ws.append([r["platform"], r["outlet_id"], r["outlet"], r["product"],
                   r["last_month_amount"], r["carried_ratio"], f"=$B$4*F{row}",
                   r["roas_last_month"], r["reason"]])
    for r in new:
        ws.append([r["platform"], r["outlet_id"], r["outlet"], r["product"], None, None,
                   None, r["roas_last_month"], r["reason"]])
        ws.cell(row=ws.max_row, column=7).fill = INPUT_FILL
    ws.append(["Total", None, None, None, f"=SUM(E{first}:E{max(last_new, last_carried)})",
               None, f"=SUM(G{first}:G{max(last_new, last_carried)})"])
    for cell in ws[ws.max_row]:
        cell.font = BOLD

    for row in ws.iter_rows(min_row=first):
        row[5].number_format = "0.0%"
        row[4].number_format = row[6].number_format = "#,##0"
        row[7].number_format = "0.00"
    for col, width in zip("ABCDEFGHI", [9, 12, 18, 8, 16, 13, 16, 15, 60]):
        ws.column_dimensions[col].width = width
    for c in ("B2", "B3", "B4", "B5", "B7"):
        ws[c].number_format = "#,##0"
    ws.column_dimensions["A"].width = 30

    cf = wb.create_sheet("carry_forward")
    cf.append(["platform", "outlet_id", "product", "amount"])
    for b in last_budget:
        cf.append([b["platform"], b["outlet_id"], b["product"], to_decimal(b["amount"])])

    dr = wb.create_sheet("dropped")
    dr.append(["platform", "outlet_id", "product", "amount", "reason"])
    for d in dropped:
        dr.append([d["platform"], d["outlet_id"], d["product"], to_decimal(d["amount"]), d["reason"]])
    for sheet in (cf, dr):
        for cell in sheet[1]:
            cell.font = BOLD
            cell.fill = HEADER_FILL

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def _read_csv(path):
    with open(path, newline="") as f:
        return list(csv.DictReader(f))


if __name__ == "__main__":
    if len(sys.argv) != 6:
        sys.exit(__doc__)
    client, month, budget_csv, funnel_csv, out = sys.argv[1:]
    config, rules = load_client(client)
    last_budget, funnel = _read_csv(budget_csv), _read_csv(funnel_csv)
    carried, new, dropped, gmv = plan(config, rules, last_budget, funnel)
    print(write_workbook(client, month, rules, carried, new, dropped, gmv, last_budget, out))
